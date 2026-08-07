import csv
import logging
from dataclasses import dataclass, field

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.models.document import DocumentType

logger = logging.getLogger(__name__)


@dataclass
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass
class ExtractionResult:
    pages: list[ExtractedPage] = field(default_factory=list)
    page_count: int | None = None

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class ExtractionError(Exception):
    pass


def extract_text(file_path: str, file_type: DocumentType) -> ExtractionResult:
    try:
        if file_type == DocumentType.PDF:
            return _extract_pdf(file_path)
        if file_type == DocumentType.DOCX:
            return _extract_docx(file_path)
        if file_type == DocumentType.TXT:
            return _extract_txt(file_path)
        if file_type == DocumentType.CSV:
            return _extract_csv(file_path)
        raise ExtractionError(f"Unsupported file type: {file_type}")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - convert any parser failure to our error type
        logger.exception("Text extraction failed for %s", file_path)
        raise ExtractionError(f"Failed to extract text: {exc}") from exc


def _extract_pdf(file_path: str) -> ExtractionResult:
    pages: list[ExtractedPage] = []
    with fitz.open(file_path) as pdf:
        for i, page in enumerate(pdf):
            text = page.get_text("text")
            pages.append(ExtractedPage(page_number=i + 1, text=text))
    return ExtractionResult(pages=pages, page_count=len(pages))


def _extract_docx(file_path: str) -> ExtractionResult:
    doc = DocxDocument(file_path)
    # DOCX has no native "pages" concept via python-docx, so we treat the
    # whole document as a single logical page for citation purposes.
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables, since policy/finance docs often live there.
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                table_text.append(row_text)

    full_text = "\n".join(paragraphs + table_text)
    return ExtractionResult(pages=[ExtractedPage(page_number=None, text=full_text)], page_count=1)


def _extract_txt(file_path: str) -> ExtractionResult:
    with open(file_path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return ExtractionResult(pages=[ExtractedPage(page_number=None, text=text)], page_count=1)


def _extract_csv(file_path: str) -> ExtractionResult:
    lines = []
    with open(file_path, encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if header:
            lines.append(" | ".join(header))
            for row in reader:
                lines.append(" | ".join(row))
    text = "\n".join(lines)
    return ExtractionResult(pages=[ExtractedPage(page_number=None, text=text)], page_count=1)
